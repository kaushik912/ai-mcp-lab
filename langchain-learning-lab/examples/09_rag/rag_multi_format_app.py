"""Retrieval-Augmented Generation over a directory of mixed file types (.html,
.docx, .md, .txt), as a LangGraph: retrieve -> generate. Same graph shape as
rag_basic_app.py; the difference here is ingesting a whole directory of
heterogeneous formats into one vector store instead of a single .txt/.pdf.
Embeddings persist to --persist-dir (Chroma on disk); a second run reuses them
instead of re-embedding - delete that directory to force a fresh embed.
Run with: python examples/09_rag/rag_multi_format_app.py [--provider {gemini,openai,openrouter}] [--source-dir <dir>] [--persist-dir <dir>] [--query <text>] [--interactive]
"""
import argparse
from pathlib import Path
from typing import TypedDict

from bs4 import BeautifulSoup
from docx import Document as DocxDocument
import markdown
from langchain_chroma import Chroma
from langchain_core.documents import Document
from langchain_core.prompts import ChatPromptTemplate
from langchain_text_splitters import RecursiveCharacterTextSplitter
from langgraph.graph import END, START, StateGraph

from common.env import data_path
from common.models import get_chat_model, get_embeddings_model
from common.prompts.rag import RAG_SYSTEM_PROMPT


def extract_text(path: Path) -> str | None:
    suffix = path.suffix.lower()
    if suffix == ".html":
        soup = BeautifulSoup(path.read_text(encoding="utf-8"), "html.parser")
        for tag in soup(["script", "style"]):
            tag.decompose()
        return soup.get_text()
    if suffix == ".docx":
        return "\n".join(p.text for p in DocxDocument(path).paragraphs)
    if suffix == ".md":
        html = markdown.markdown(path.read_text(encoding="utf-8"))
        return BeautifulSoup(html, "html.parser").get_text()
    if suffix == ".txt":
        return path.read_text(encoding="utf-8")
    return None


def build_models(provider: str):
    return get_chat_model(provider), get_embeddings_model(provider)


def load_and_chunk(source_dir: str):
    documents = []
    for path in sorted(Path(source_dir).iterdir()):
        text = extract_text(path)
        if text is None:
            print(f"Unsupported file type: {path.name}")
            continue
        documents.append(Document(page_content=text, metadata={"source": path.name}))

    text_splitter = RecursiveCharacterTextSplitter(chunk_size=400, chunk_overlap=100)
    return text_splitter.split_documents(documents)


def get_vector_store(source_dir: str, persist_dir: str, embeddings):
    if Path(persist_dir).exists():
        print(f"Loading existing embeddings from {persist_dir}")
        return Chroma(persist_directory=persist_dir, embedding_function=embeddings)

    print(f"No embeddings found at {persist_dir}, embedding {source_dir}")
    chunks = load_and_chunk(source_dir)
    return Chroma.from_documents(chunks, embeddings, persist_directory=persist_dir)


class RAGState(TypedDict):
    query: str
    context: list[Document]
    answer: str


def build_graph(llm, vector_store):
    prompt_template = ChatPromptTemplate.from_messages(
        [("system", RAG_SYSTEM_PROMPT), ("human", "{question}")]
    )

    def retrieve(state: RAGState) -> dict:
        return {"context": vector_store.similarity_search(state["query"], k=3)}

    def generate(state: RAGState) -> dict:
        context_text = "\n\n".join(doc.page_content for doc in state["context"])
        prompt = prompt_template.invoke({"context": context_text, "question": state["query"]})
        response = llm.invoke(prompt)
        return {"answer": response.text}

    builder = StateGraph(RAGState)
    builder.add_node("retrieve", retrieve)
    builder.add_node("generate", generate)
    builder.add_edge(START, "retrieve")
    builder.add_edge("retrieve", "generate")
    builder.add_edge("generate", END)
    return builder.compile()


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--provider", choices=["gemini", "openai", "openrouter"], default="openrouter")
    parser.add_argument("--source-dir", default=str(data_path("rag-multi")))
    parser.add_argument("--persist-dir", default=str(data_path("rag-multi-chroma")), help="delete this directory to force re-embedding")
    parser.add_argument("--query", default=None)
    parser.add_argument("--interactive", action="store_true", help="loop for multiple questions against one indexed session")
    args = parser.parse_args()

    llm, embeddings = build_models(args.provider)
    vector_store = get_vector_store(args.source_dir, args.persist_dir, embeddings)

    graph = build_graph(llm, vector_store)

    if args.interactive:
        print("Interactive mode. Type 'quit' or 'exit' to stop.")
        while True:
            query = input("\nEnter your question: ").strip()
            if query.lower() in ("quit", "exit"):
                break
            if not query:
                continue
            result = graph.invoke({"query": query, "context": [], "answer": ""})
            print(result["answer"])
        return

    query = args.query or input("Enter your question: ")
    result = graph.invoke({"query": query, "context": [], "answer": ""})
    print(result["answer"])


if __name__ == "__main__":
    main()
